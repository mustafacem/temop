import streamlit as st
from proposal_droid.chat_gpt.chat_gpt import  ask_chatgpt , checker
from proposal_droid.doc_creation.doc_creation import create_excel_with_values_2
from proposal_droid.ocr_processing.ocr_processing import preprocess_handwritten_image, extract_text_from_image
from proposal_droid.whisper_speech_to_text.whisper_speech_to_text import transcribe_audio # transcribe_czech_audio, transcribe_english_audio
from proposal_droid.data_from_web.data_from_web import transcribe_english_youtube, extract_text_from_url
from kd_streamlit import st_init, LanguageUI
import os
from io import BytesIO
import openai
from docx import Document
from dotenv import load_dotenv

def create_docx(items_dict):
    doc = Document()
    doc.add_heading('Use Case Description', 0)
    if 'Use Case Description' in items_dict:
        doc.add_paragraph(items_dict['Use Case Description'])

    for key, value in items_dict.items():
        if key != 'Use Case Description':
            doc.add_heading(key, level=1)
            doc.add_paragraph(value)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def main():
    
    openai_api_key = st.sidebar.text_input("Enter your OpenAI API Key", type="password")
    if openai_api_key:
        openai.api_key = openai_api_key

    # Initialize session state variables
    if "notes" not in st.session_state:
        st.session_state.notes = None

    if "items_dict" not in st.session_state:
        st.session_state.items_dict = {
            "Use Case Description": None,
            "Target functionality": None,
            "Solution": None,
            "Inputs": None,
            "Outputs": None,
            "Requirements and assumptions": None,
            "PoC vs production implementation": None,
            "Human review of outputs": None,
            "Post-PoC improvements and functionality add-ons": None,
            "Infrastructure": None,
            "Limitations": None
        }

    if "current_item_index" not in st.session_state:
        st.session_state.current_item_index = 0

    if "notes_finalized" not in st.session_state:
        st.session_state.notes_finalized = False

    if "use_case_finalized" not in st.session_state:
        st.session_state.use_case_finalized = False

    if "use_case_description" not in st.session_state:
        st.session_state.use_case_description = ""

    # Use tabs to organize the workflow
    tabs = ["Upload Notes", "Edit and Finalize Notes", "Generate Use Case", "Proposal Items", "Download"]

    # Determine the active tab based on session state
    if not st.session_state.notes:
        active_tab = "Upload Notes"
    elif not st.session_state.notes_finalized:
        active_tab = "Edit and Finalize Notes"
    elif not st.session_state.use_case_finalized:
        active_tab = "Generate Use Case"
    elif st.session_state.current_item_index < len(st.session_state.items_dict.keys()) - 1:  # Exclude 'Use Case Description'
        active_tab = "Proposal Items"
    else:
        active_tab = "Download"

    # Display tabs
    selected_tab = st.sidebar.radio("Navigation", tabs, index=tabs.index(active_tab))

    # Tab 1: Upload Notes
    if selected_tab == "Upload Notes":
        file_type = st.selectbox("Select the type of notes you want to upload:", ["Image", "Voice Recording", "YouTube Link", "Website Link"])
        uploaded_file = None

        # Handling different file types
        if file_type in ["Image", "Voice Recording"]:
            uploaded_file = st.file_uploader("Upload your file", type=["png", "jpg", "jpeg", "wav", "mp3", "m4a"])

        if file_type == "Image" and uploaded_file is not None:
            st.image(uploaded_file, caption="Uploaded Image", use_column_width=True)
            preprocess = st.checkbox("Preprocess Image")
            preprocessed_image = uploaded_file

            if preprocess:
                preprocessed_image = preprocess_handwritten_image(uploaded_file)
                if preprocessed_image is not None:
                    st.image(preprocessed_image, caption="Preprocessed Image", use_column_width=True)

            # ocr_method = st.selectbox("Select OCR method:", ["Method 1", "Method 2", "Method 3"])
            if st.button("Extract Text from Image"):
                image_to_process = preprocessed_image if preprocess else uploaded_file
                st.session_state.notes = extract_text_from_image(openai_api_key, image_to_process)
                st.success("Text extracted from image!")
                st.session_state.notes_finalized = False  # Reset for navigation

        elif file_type == "Voice Recording" and uploaded_file is not None:
            st.audio(uploaded_file, format="audio/wav" if uploaded_file.name.endswith('.wav') else "audio/mp3")
            language = st.radio("Select the language of the audio:", ("Czech", "English"))

            if st.button("Transcribe Audio"):
                with st.spinner('Transcribing audio...'):
                    if language == "Czech":
                        st.session_state.notes = transcribe_audio(uploaded_file,"czech",15)
                    elif language == "English":
                        st.session_state.notes = transcribe_audio(uploaded_file,"english",15)
                st.success("Transcription completed!")
                st.session_state.notes_finalized = False  # Reset for navigation

        elif file_type == "YouTube Link":
            youtube_link = st.text_input("Enter YouTube link:")
            if youtube_link and st.button("Process YouTube Link"):
                with st.spinner('Processing YouTube link...'):
                    st.session_state.notes = transcribe_english_youtube(youtube_link)
                st.success("YouTube video processed!")
                st.session_state.notes_finalized = False  # Reset for navigation

        elif file_type == "Website Link":
            website_link = st.text_input("Enter website link:")
            if website_link and st.button("Extract Text from Website"):
                with st.spinner('Extracting text...'):
                    st.session_state.notes = extract_text_from_url(website_link)
                st.success("Website text extracted!")
                st.session_state.notes_finalized = False  # Reset for navigation

        # If notes are available, allow navigation to the next tab
        if st.session_state.notes:
            st.sidebar.success("Notes are ready! Proceed to 'Edit and Finalize Notes' tab.")

    # Tab 2: Edit and Finalize Notes
    elif selected_tab == "Edit and Finalize Notes":
        if "notes" in st.session_state and st.session_state.notes:
            st.write("Transcribed Notes:")
            st.write(st.session_state.notes)

            if "edit_mode" not in st.session_state:
                st.session_state.edit_mode = False

            if st.button("Edit Notes"):
                st.session_state.edit_mode = True

            if st.session_state.edit_mode:
                edited_notes = st.text_area("Edit your notes here:", st.session_state.notes, key="edit_notes_input")
                if st.button("Save Edited Notes"):
                    st.session_state.notes = edited_notes
                    st.session_state.edit_mode = False
                    st.success("Notes updated!")

            if st.button("Use Notes As Is"):
                st.session_state.notes_finalized = True
                st.success("Notes finalized!")
        else:
            st.warning("Please upload and extract notes in the 'Upload Notes' tab.")

    # Tab 3: Generate Use Case
    elif selected_tab == "Generate Use Case":
        if st.session_state.notes_finalized:
            if not st.session_state.use_case_description:
                st.session_state.use_case_description = ask_chatgpt(
                    "Generate a very short use case description from the given notes: " + st.session_state.notes
                )

            st.subheader("Use Case Description")
            st.write(st.session_state.use_case_description)
            edited_description = st.text_area(
                "Edit Use Case Description:", st.session_state.use_case_description, key="use_case_desc_input"
            )
            if st.button("Finalize Use Case Description"):
                st.session_state.use_case_description = edited_description
                st.session_state.use_case_finalized = True
                st.session_state.items_dict["Use Case Description"] = st.session_state.use_case_description
                st.success("Use case description finalized!")
        else:
            st.warning("Please finalize your notes in the 'Edit and Finalize Notes' tab.")

    # Tab 4: Proposal Items
    elif selected_tab == "Proposal Items":
        if st.session_state.use_case_finalized:
            st.subheader("Use Case Description")
            st.write(st.session_state.use_case_description)
            if st.button("Edit Use Case Description"):
                st.session_state.use_case_finalized = False
                # Go back to Generate Use Case tab
                selected_tab = "Generate Use Case"
                st.experimental_rerun()

            keys = list(st.session_state.items_dict.keys())
            # Exclude 'Use Case Description' from proposal items
            proposal_keys = [k for k in keys if k != "Use Case Description"]

            if st.session_state.current_item_index < len(proposal_keys):
                current_key = proposal_keys[st.session_state.current_item_index]

                if st.session_state.items_dict[current_key] is None:
                    prompt = f"Generate a short 2-4 sentence {current_key} for a business proposal based on the following use case description: {st.session_state.use_case_description}"
                    value = ask_chatgpt(prompt)
                    st.session_state.items_dict[current_key] = value
                    opinion_ai = checker(st.session_state.items_dict[current_key], value)
                else:
                    opinion_ai = ""

                st.write(f"**{current_key}:**")
                st.write(st.session_state.items_dict[current_key])
                st.write(f"**AI Opinion for {current_key}:**\n{opinion_ai}")

                action = st.selectbox(
                    f"Select action for {current_key}:", ["Accept", "Edit", "Add"], key=f"action_select_{current_key}"
                )

                if action == "Accept" and st.button(f"Confirm {current_key}", key=f"confirm_{current_key}"):
                    st.session_state.current_item_index += 1
                    st.success(f"{current_key} accepted!")

                elif action == "Edit":
                    edited_value = st.text_area(
                        f"Edit {current_key}:", st.session_state.items_dict[current_key], key=f"edit_{current_key}"
                    )
                    if st.button(f"Save {current_key}", key=f"save_{current_key}"):
                        st.session_state.items_dict[current_key] = edited_value
                        #st.session_state.current_item_index += 1
                        st.success(f"{current_key} updated!")
                        opinion_ai = checker(st.session_state.items_dict[current_key], edited_value)

                elif action == "Add":
                    additional_value = st.text_area(f"Add to {current_key}:", key=f"add_{current_key}")
                    if st.button(f"Add to {current_key}", key=f"add_confirm_{current_key}"):
                        # Create a new prompt for adding content based on the additional input
                        prompt_for_reg = f"Expand the {current_key} for the business proposal based on the following additional information: {additional_value} and the use case description: {st.session_state.use_case_description}"
                        
                        # Ask ChatGPT to generate new content based on the additional information
                        new_value = ask_chatgpt(prompt_for_reg)
                        
                        # Append the new value to the existing content (you can modify how you combine this)
                        st.session_state.items_dict[current_key] += f"\nAdditional content:\n{new_value}"
                        
                        # Move to the next item
                        #st.session_state.current_item_index += 1
                        st.success(f"Added to {current_key}!")


            else:
                st.success("All items processed! Proceed to the 'Download' tab.")
        else:
            st.warning("Please finalize the use case description in the 'Generate Use Case' tab.")

    # Tab 5: Download
    elif selected_tab == "Download":
        if st.session_state.current_item_index >= len(st.session_state.items_dict.keys()) - 1:
            # Create and download the DOCX file
            docx_buffer = create_docx(st.session_state.items_dict)
            st.download_button(
                label="Download DOCX",
                data=docx_buffer,
                file_name='output.docx',
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

            # Create and download the Excel file
            excel_buffer = create_excel_with_values_2(st.session_state.notes)  # Ensure this function exists
            st.download_button(
                label="Download Excel",
                data=excel_buffer,
                file_name='price_estimate.xlsx',
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            st.warning("Please complete all proposal items in the 'Proposal Items' tab.")

    # Reset button at the sidebar
    if st.sidebar.button("Reset"):
        st.session_state.clear()
        st.success("Reset completed!")

if __name__ == "__main__":
    load_dotenv(override=True)

    st_init(auth=True, notice=False, feedback=False, title="ProposalDroid", language=LanguageUI.CZ)
    main()
